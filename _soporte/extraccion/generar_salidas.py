"""Genera la planilla de revisión humana (revision/<area>_planilla.xlsx) a partir de los JSON verificados.
La transcripción legible la publica extraccion/publicar.py en DBA/ y ESTANDARES/.

    python3 extraccion/generar_salidas.py matematicas
"""
import csv
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


NOMBRES = {'matematicas': 'Matemáticas', 'lenguaje': 'Lenguaje', 'naturales': 'Ciencias Naturales',
           'sociales': 'Ciencias Sociales', 'ingles': 'Inglés'}


def main(area):
    base = f'_soporte/datos/{area}'
    est = json.load(open(f'{base}/estandares.json', encoding='utf-8'))
    dba = json.load(open(f'{base}/dba.json', encoding='utf-8'))

    os.makedirs('_soporte/revision', exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Revisión'
    cab = ['Página PDF', 'Documento', 'Tipo', 'Código', 'Texto extraído', '¿Coincide con la página?', 'Corrección o comentario', 'Revisó', 'Fecha']
    ws.append(cab)
    filas = []
    for e in est:
        filas.append((e['pagina_pdf'], e['fuente'], 'Estándar', e['codigo'], e['texto']))
        filas += [(e['pagina_pdf'], e['fuente'], 'Subproceso', sp['codigo'], sp['texto']) for sp in e.get('subprocesos', [])]
    for d in dba:
        filas.append((d['pagina_pdf'], d['fuente'], 'DBA', d['codigo'], d['enunciado']))
        filas += [(d['pagina_pdf'], d['fuente'], 'Evidencia', e['codigo'], e['texto']) for e in d['evidencias']]
    for f in filas:
        ws.append(list(f) + ['', '', '', ''])
    anchos = [10, 26, 11, 22, 90, 16, 40, 18, 12]
    for i, a in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = a
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3A5F')
    for fila in ws.iter_rows(min_row=2):
        fila[4].alignment = Alignment(wrap_text=True, vertical='top')
    dv = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'F2:F{len(filas) + 1}')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:I{len(filas) + 1}'
    ins = wb.create_sheet('Instrucciones')
    for linea in [
        'Planilla de revisión humana de la extracción.',
        '1. Abra el PDF indicado en la columna Documento, en la página de la columna Página PDF.',
        '2. Compare el texto extraído con el de la página: mismas palabras, mismo orden, mismo corte (que no le falte ni le sobre nada).',
        '3. Marque Sí o No. Si es No, escriba en Corrección exactamente qué dice el PDF.',
        '4. Escriba su nombre y la fecha.',
        'La máquina ya verificó que cada texto aparece literal en el PDF; esta revisión confirma que cada texto quedó bien separado y en el código correcto.',
    ]:
        ins.append([linea])
    ins.column_dimensions['A'].width = 130
    wb.save(f'_soporte/revision/{area}_planilla.xlsx')
    print(len(est), 'estándares,', sum(len(e.get('subprocesos', [])) for e in est), 'subprocesos,', len(dba), 'DBA,', sum(len(d['evidencias']) for d in dba), 'evidencias → planilla')


if __name__ == '__main__':
    main(sys.argv[1])
