"""Genera, a partir de los JSON verificados, las salidas para otros usos. Nunca se editan a mano:
si algo cambia, se corrige la extracción y se vuelven a generar.

- datos/<area>/estandares.csv, dba.csv, evidencias.csv  (hojas de cálculo, sistemas)
- datos/<area>/<area>.md                                (lectura)
- revision/<area>_planilla.xlsx                          (revisión humana página por página)

    python3 extraccion/generar_salidas.py matematicas
"""
import csv
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


NOMBRES = {'matematicas': 'Matemáticas', 'lenguaje': 'Lenguaje', 'naturales': 'Ciencias Naturales',
           'sociales': 'Ciencias Sociales', 'ingles': 'Inglés'}


def main(area):
    base = f'datos/{area}'
    est = json.load(open(f'{base}/estandares.json', encoding='utf-8'))
    dba = json.load(open(f'{base}/dba.json', encoding='utf-8'))

    with open(f'{base}/estandares.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['codigo', 'grado_desde', 'grado_hasta', 'eje', 'eje_nombre', 'orden', 'texto', 'fuente', 'pagina_pdf'])
        for e in est:
            w.writerow([e['codigo'], e['grado_desde'], e['grado_hasta'], e['eje'], e['eje_nombre'], e['orden'], e['texto'], e['fuente'], e['pagina_pdf']])
    if any(e.get('subprocesos') for e in est):
        with open(f'{base}/subprocesos.csv', 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(['codigo', 'estandar', 'grado_desde', 'grado_hasta', 'numero', 'texto', 'fuente', 'pagina_pdf'])
            for e in est:
                for sp in e.get('subprocesos', []):
                    w.writerow([sp['codigo'], e['codigo'], e['grado_desde'], e['grado_hasta'], sp['numero'], sp['texto'], e['fuente'], e['pagina_pdf']])
    with open(f'{base}/dba.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['codigo', 'grado', 'numero', 'enunciado', 'fuente', 'pagina_pdf'])
        for d in dba:
            w.writerow([d['codigo'], d['grado'], d['numero'], d['enunciado'], d['fuente'], d['pagina_pdf']])
    with open(f'{base}/evidencias.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['codigo', 'dba', 'grado', 'numero', 'texto', 'fuente', 'pagina_pdf'])
        for d in dba:
            for e in d['evidencias']:
                w.writerow([e['codigo'], d['codigo'], d['grado'], e['numero'], e['texto'], d['fuente'], d['pagina_pdf']])

    L = [f'# {NOMBRES.get(area, area)}: Estándares Básicos de Competencias y Derechos Básicos de Aprendizaje', '',
         'Texto extraído y verificado contra los PDF oficiales del MEN (ver `verificacion/`). '
         'Archivo generado: no editar a mano.', '', '## Estándares Básicos de Competencias', '']
    ciclo = None
    for e in est:
        c = (e['grado_desde'], e['grado_hasta'])
        if c != ciclo:
            ciclo, eje = c, None
            L += [f'### Grados {c[0]}° a {c[1]}°', '']
        if e['eje_nombre'] != eje:
            eje = e['eje_nombre']
            L += [f'**{eje.capitalize()}**', '']
        if e.get('subprocesos'):
            L += [f"- **`{e['codigo']}`** {e['texto']}", '  *Para lo cual,*']
            L += [f"  - `{sp['codigo']}` {sp['texto']}" for sp in e['subprocesos']]
        else:
            L += [f"- `{e['codigo']}` {e['texto']}"]
        if est.index(e) + 1 < len(est) and est[est.index(e) + 1]['eje_nombre'] != eje:
            L += ['']
    L += ['', '## Derechos Básicos de Aprendizaje', '']
    grado = None
    for d in dba:
        if d['grado'] != grado:
            grado = d['grado']
            L += [f'### Grado {grado}°', '']
        L += [f"**`{d['codigo']}`** {d['enunciado']}", '', '*Evidencias de aprendizaje*', '']
        L += [f"- `{e['codigo']}` {e['texto']}" for e in d['evidencias']] + ['']
    with open(f'{base}/{area}.md', 'w', encoding='utf-8') as f:
        f.write('\n'.join(L) + '\n')

    os.makedirs('revision', exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Revisión'
    cab = ['Página PDF', 'Documento', 'Tipo', 'Código', 'Texto extraído', '¿Coincide con la página?', 'Corrección o comentario', 'Revisó', 'Fecha']
    ws.append(cab)
    filas = []
    for e in est:
        filas.append((e['pagina_pdf'], e['fuente'], 'Estándar', e['codigo'], e['texto']))
        filas += [(e['pagina_pdf'], e['fuente'], 'Subproceso', sp['codigo'], sp['texto']) for sp in e.get('subprocesos', [])]
    for d in dba:
        filas.append((d['pagina_pdf'], d['fuente'], 'DBA', d['codigo'], d['enunciado']))
        filas += [(d['pagina_pdf'], d['fuente'], 'Evidencia', e['codigo'], e['texto']) for e in d['evidencias']]
    for f in filas:
        ws.append(list(f) + ['', '', '', ''])
    anchos = [10, 26, 11, 22, 90, 16, 40, 18, 12]
    for i, a in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = a
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F3A5F')
    for fila in ws.iter_rows(min_row=2):
        fila[4].alignment = Alignment(wrap_text=True, vertical='top')
    dv = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'F2:F{len(filas) + 1}')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:I{len(filas) + 1}'
    ins = wb.create_sheet('Instrucciones')
    for linea in [
        'Planilla de revisión humana de la extracción.',
        '1. Abra el PDF indicado en la columna Documento, en la página de la columna Página PDF.',
        '2. Compare el texto extraído con el de la página: mismas palabras, mismo orden, mismo corte (que no le falte ni le sobre nada).',
        '3. Marque Sí o No. Si es No, escriba en Corrección exactamente qué dice el PDF.',
        '4. Escriba su nombre y la fecha.',
        'La máquina ya verificó que cada texto aparece literal en el PDF; esta revisión confirma que cada texto quedó bien separado y en el código correcto.',
    ]:
        ins.append([linea])
    ins.column_dimensions['A'].width = 130
    wb.save(f'revision/{area}_planilla.xlsx')
    print(len(est), 'estándares,', sum(len(e.get('subprocesos', [])) for e in est), 'subprocesos,', len(dba), 'DBA,', sum(len(d['evidencias']) for d in dba), 'evidencias → CSV, MD y planilla')


if __name__ == '__main__':
    main(sys.argv[1])
